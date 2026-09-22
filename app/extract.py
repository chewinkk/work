"""Pull readable text out of the files a course hands you.

Shared by the deck generator and the assignment assistant, so a source
document and an attached specification are read the same way.
"""
import base64
import os
from html.parser import HTMLParser

MAX_SOURCE_CHARS = 600_000   # ~150k tokens, well inside the 1M context window
MAX_PDF_BYTES = 25 * 1024 * 1024

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json"}
READABLE_SUFFIXES = TEXT_SUFFIXES | {".pdf", ".docx", ".pptx"}

_SKIP_TAGS = {"script", "style", "head"}
_BREAK_TAGS = {"p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}


class _Stripper(HTMLParser):
    """Canvas stores assignment descriptions as HTML. Claude wants text."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in _SKIP_TAGS:
            self._skip += 1
        elif tag in _BREAK_TAGS:
            self.parts.append("\n")
        elif tag == "li":
            self.parts.append("\n- ")

    def handle_endtag(self, tag):
        if tag in _SKIP_TAGS and self._skip:
            self._skip -= 1
        elif tag in _BREAK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)

    def text(self):
        joined = "".join(self.parts)
        lines = [line.strip() for line in joined.splitlines()]
        return "\n".join(line for line in lines if line)


def html_to_text(html):
    if not html:
        return ""
    parser = _Stripper()
    try:
        parser.feed(html)
        parser.close()
    except Exception:
        return html
    return parser.text()


def extract_docx(path):
    from docx import Document

    document = Document(path)
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_pptx(path):
    from pptx import Presentation

    parts = []
    for slide in Presentation(path).slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
    return "\n".join(parts)


def extract_pdf(path):
    """Text only. The deck generator sends PDFs to Claude natively instead."""
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    try:
        return "\n".join((page.extract_text() or "") for page in PdfReader(path).pages)
    except Exception:
        return ""


def extract_text(path, filename=None):
    """Best effort plain text. Returns '' when the type is unreadable."""
    suffix = os.path.splitext(filename or path)[1].lower()
    try:
        if suffix in TEXT_SUFFIXES:
            with open(path, "r", encoding="utf-8", errors="replace") as handle:
                return handle.read().strip()
        if suffix == ".docx":
            return extract_docx(path).strip()
        if suffix == ".pptx":
            return extract_pptx(path).strip()
        if suffix == ".pdf":
            return extract_pdf(path).strip()
    except Exception:
        return ""
    return ""


def source_block(path, filename):
    """A Claude content block for an uploaded source document.

    PDFs go over as native document blocks so tables and layout survive.
    Everything else is extracted to text first.
    """
    suffix = os.path.splitext(filename)[1].lower()

    if suffix == ".pdf":
        size = os.path.getsize(path)
        if size > MAX_PDF_BYTES:
            raise RuntimeError(
                f"{filename} is {size // 1024 // 1024} MB. "
                "Source PDFs must be under 25 MB to send to Claude."
            )
        with open(path, "rb") as handle:
            data = base64.standard_b64encode(handle.read()).decode("ascii")
        return {
            "type": "document",
            "source": {"type": "base64", "media_type": "application/pdf", "data": data},
        }

    if suffix not in READABLE_SUFFIXES:
        raise RuntimeError(
            f"Cannot read {suffix or 'that file type'}. Supported source files: "
            "PDF, DOCX, PPTX, TXT, MD, CSV."
        )

    text = extract_text(path, filename)
    if not text:
        raise RuntimeError(f"No readable text found in {filename}.")
    if len(text) > MAX_SOURCE_CHARS:
        raise RuntimeError(
            f"{filename} holds {len(text):,} characters, over the "
            f"{MAX_SOURCE_CHARS:,} limit. Split the document and try again."
        )
    return {"type": "text", "text": text}
